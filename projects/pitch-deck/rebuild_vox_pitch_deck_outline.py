from __future__ import annotations

import argparse
from dataclasses import dataclass
from io import BytesIO
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE, MSO_CONNECTOR_TYPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


DEFAULT_SOURCE_DECK = Path("input/Vox PD v2.0.pptx")
DEFAULT_OUTPUT_DECK = Path("output/pitch-deck/vox-pd-v3-outline-stub.pptx")


BLACK = RGBColor(0x00, 0x00, 0x00)
ACCENT = RGBColor(0x51, 0x83, 0x94)
ACCENT_DARK = RGBColor(0x3B, 0x67, 0x77)
ACCENT_LIGHT = RGBColor(0xE9, 0xF1, 0xF4)
WARM_PANEL = RGBColor(0xF7, 0xF3, 0xEC)
SOFT_PANEL = RGBColor(0xFB, 0xF9, 0xF5)
MID_GRAY = RGBColor(0xD5, 0xD3, 0xCF)
LIGHT_GRAY = RGBColor(0xEF, 0xEE, 0xEB)
LINE = RGBColor(0xBF, 0xB8, 0xAF)
TEXT_SOFT = RGBColor(0x44, 0x40, 0x3C)
TEXT_MUTED = RGBColor(0x51, 0x83, 0x94)


@dataclass
class ImageAsset:
    blob: bytes
    left: int
    top: int
    width: int
    height: int


def emu(inches: float) -> int:
    return Inches(inches)


def pt(points: float) -> int:
    return Pt(points)


def delete_slides_after(prs: Presentation, keep_count: int) -> None:
    sld_id_lst = prs.slides._sldIdLst
    for sld_id in list(sld_id_lst)[keep_count:]:
        prs.part.drop_rel(sld_id.rId)
        sld_id_lst.remove(sld_id)


def clear_slide(slide) -> None:
    sp_tree = slide.shapes._spTree
    for shape in list(slide.shapes):
        sp_tree.remove(shape._element)


def add_textbox(
    slide,
    x: float,
    y: float,
    w: float,
    h: float,
    text: str,
    *,
    font_name: str = "Montserrat",
    font_size: float = 18,
    color: RGBColor = BLACK,
    bold: bool | None = None,
    align: PP_ALIGN = PP_ALIGN.LEFT,
    vertical_anchor: MSO_ANCHOR = MSO_ANCHOR.TOP,
    margins: tuple[float, float, float, float] = (0.0, 0.0, 0.0, 0.0),
    line_spacing: float | None = None,
) :
    box = slide.shapes.add_textbox(emu(x), emu(y), emu(w), emu(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = vertical_anchor
    tf.margin_left = emu(margins[0])
    tf.margin_top = emu(margins[1])
    tf.margin_right = emu(margins[2])
    tf.margin_bottom = emu(margins[3])
    p = tf.paragraphs[0]
    p.alignment = align
    if line_spacing is not None:
        p.line_spacing = pt(line_spacing)
    run = p.add_run()
    run.text = text
    run.font.name = font_name
    run.font.size = pt(font_size)
    run.font.color.rgb = color
    if bold is not None:
        run.font.bold = bold
    return box


def add_box(
    slide,
    x: float,
    y: float,
    w: float,
    h: float,
    *,
    fill: RGBColor = SOFT_PANEL,
    line: RGBColor = LINE,
    line_width: float = 1.0,
    shape_type=MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE,
):
    shape = slide.shapes.add_shape(shape_type, emu(x), emu(y), emu(w), emu(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = line
    shape.line.width = pt(line_width)
    return shape


def add_line(slide, x1: float, y1: float, x2: float, y2: float, *, color: RGBColor = BLACK, width: float = 0.7):
    line = slide.shapes.add_connector(
        MSO_CONNECTOR_TYPE.STRAIGHT,
        emu(x1),
        emu(y1),
        emu(x2),
        emu(y2),
    )
    line.line.color.rgb = color
    line.line.width = pt(width)
    return line


def add_chip(slide, x: float, y: float, w: float, h: float, text: str) -> None:
    chip = add_box(slide, x, y, w, h, fill=SOFT_PANEL, line=LINE, line_width=0.8)
    tf = chip.text_frame
    tf.clear()
    tf.word_wrap = False
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = emu(0.12)
    tf.margin_right = emu(0.12)
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = text
    run.font.name = "Inter"
    run.font.size = pt(8.6)
    run.font.color.rgb = TEXT_SOFT


def add_footer(slide, footer_bird: bytes) -> None:
    add_line(slide, 0.583, 5.167, 8.58, 5.167, color=BLACK, width=0.72)
    slide.shapes.add_picture(BytesIO(footer_bird), emu(8.58), emu(4.42), width=emu(1.052), height=emu(1.052))


def add_header(slide, title: str, subtitle: str | None = None) -> None:
    add_textbox(slide, 0.58, 0.33, 8.75, 0.44, title, font_name="Montserrat", font_size=18.2, color=BLACK, bold=False)
    if subtitle:
        add_textbox(
            slide,
            0.58,
            0.82,
            8.15,
            0.42,
            subtitle,
            font_name="Inter",
            font_size=10.6,
            color=TEXT_MUTED,
            bold=False,
        )


def add_centered_header(slide, title: str, subtitle: str | None = None) -> None:
    add_textbox(
        slide,
        0.65,
        0.34,
        8.7,
        0.45,
        title,
        font_name="Montserrat",
        font_size=21.2,
        color=BLACK,
        bold=False,
        align=PP_ALIGN.CENTER,
    )
    if subtitle:
        add_textbox(
            slide,
            1.0,
            0.82,
            8.0,
            0.38,
            subtitle,
            font_name="Inter",
            font_size=10.4,
            color=TEXT_MUTED,
            align=PP_ALIGN.CENTER,
        )


def add_panel_label(slide, x: float, y: float, text: str, *, accent: bool = False) -> None:
    color = ACCENT_DARK if accent else TEXT_SOFT
    add_textbox(slide, x, y, 2.0, 0.18, text, font_name="Inter", font_size=8.4, color=color, bold=True)


def build_slide_1(slide, footer_bird: bytes) -> None:
    clear_slide(slide)
    add_line(slide, 1.30, 0.74, 8.70, 0.74, color=BLACK, width=0.72)
    add_line(slide, 2.02, 4.94, 7.98, 4.94, color=BLACK, width=0.72)
    add_textbox(
        slide,
        4.06,
        2.40,
        1.90,
        0.40,
        "VOX",
        font_name="Montserrat",
        font_size=27.5,
        color=BLACK,
        bold=False,
        align=PP_ALIGN.CENTER,
    )
    add_textbox(
        slide,
        2.02,
        3.26,
        5.96,
        0.32,
        "Pre-visit clinical intelligence for efficient, high-quality care.",
        font_name="Montserrat",
        font_size=11.8,
        color=BLACK,
        align=PP_ALIGN.CENTER,
    )


def build_slide_2(slide, footer_bird: bytes) -> None:
    clear_slide(slide)
    add_header(slide, "Demand is outpacing clinical capacity.")

    tile_specs = [
        (0.86, "Access", "~1 month", "Average wait for a new doctor appointment", "Up ~20% in 3 years"),
        (5.06, "Deferred care", "~1 in 10", "Adults deferred care because they could not get a timely appointment", "Last year"),
    ]
    for x, label, metric, detail, kicker in tile_specs:
        tile = add_box(slide, x, 1.68, 3.62, 2.36, fill=WARM_PANEL, line=LINE, line_width=1.0)
        add_chip(slide, x + 0.22, 1.90, 1.10 if label == "Access" else 1.52, 0.32, label)
        add_textbox(slide, x + 0.22, 2.34, 2.35, 0.58, metric, font_name="Inter", font_size=30.5, color=BLACK, bold=True)
        add_textbox(slide, x + 0.22, 2.92, 3.0, 0.46, detail, font_name="Inter", font_size=10.6, color=BLACK)
        add_textbox(slide, x + 0.22, 3.56, 2.4, 0.24, kicker, font_name="Inter", font_size=8.9, color=TEXT_MUTED)
        tile.shadow.inherit = False

    add_footer(slide, footer_bird)


def build_slide_3(slide, footer_bird: bytes) -> None:
    clear_slide(slide)
    add_header(slide, "More efficient visits will increase supply.")
    add_chip(slide, 0.58, 0.84, 1.34, 0.32, "Every minute matters")
    add_chip(slide, 2.03, 0.84, 1.70, 0.32, "Small per-visit gains compound")

    add_textbox(slide, 0.82, 1.92, 1.2, 0.18, "Typical visit", font_name="Inter", font_size=8.5, color=TEXT_SOFT, bold=True)
    history = add_box(slide, 0.82, 2.18, 3.20, 0.92, fill=ACCENT, line=ACCENT, line_width=1.0)
    decisions = add_box(slide, 4.02, 2.18, 1.94, 0.92, fill=MID_GRAY, line=MID_GRAY, line_width=1.0)
    other = add_box(slide, 5.96, 2.18, 1.98, 0.92, fill=LIGHT_GRAY, line=LIGHT_GRAY, line_width=1.0)
    for box in [history, decisions, other]:
        box.line.width = pt(0.5)
    add_textbox(slide, 1.18, 2.47, 2.45, 0.28, "History reconstruction", font_name="Inter", font_size=14.8, color=RGBColor(0xFF, 0xFF, 0xFF), bold=True, align=PP_ALIGN.CENTER)
    add_textbox(slide, 4.12, 2.50, 1.72, 0.22, "Decisions + counseling", font_name="Inter", font_size=9.0, color=TEXT_SOFT, bold=True, align=PP_ALIGN.CENTER)
    add_textbox(slide, 6.14, 2.50, 1.64, 0.22, "Everything else", font_name="Inter", font_size=9.0, color=TEXT_SOFT, bold=True, align=PP_ALIGN.CENTER)

    callout = add_box(slide, 6.15, 1.56, 2.85, 1.62, fill=SOFT_PANEL, line=ACCENT, line_width=1.1)
    add_textbox(slide, 6.39, 1.82, 2.2, 0.24, "Why this matters", font_name="Montserrat", font_size=15.2, color=BLACK, bold=True)
    bullet_y = 2.22
    bullets = [
        "Most variable minutes of the visit",
        "Best target for upstream prep without losing quality",
        "Most amenable to automation",
    ]
    for bullet in bullets:
        dot = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.OVAL, emu(6.42), emu(bullet_y + 0.05), emu(0.07), emu(0.07))
        dot.fill.solid()
        dot.fill.fore_color.rgb = ACCENT
        dot.line.color.rgb = ACCENT
        add_textbox(slide, 6.55, bullet_y, 2.15, 0.22, bullet, font_name="Inter", font_size=9.6, color=BLACK, bold=True)
        bullet_y += 0.32
    add_line(slide, 4.02, 2.50, 6.15, 2.50, color=ACCENT, width=1.0)

    add_footer(slide, footer_bird)


def build_slide_4(slide, footer_bird: bytes) -> None:
    clear_slide(slide)
    add_header(
        slide,
        "Visits start with story, not structure.",
        "Clinicians spend valuable minutes turning patient story into decision-ready structure.",
    )

    story = add_box(slide, 0.78, 1.68, 2.55, 2.10, fill=WARM_PANEL, line=LINE, line_width=1.0)
    add_panel_label(slide, 1.02, 1.92, "Patient story")
    fragments = [
        '"It started a few weeks ago..."',
        '"Worse after meals..."',
        '"I am not sure if it is related..."',
        '"The pain comes and goes..."',
        '"It feels different at night..."',
    ]
    y = 2.24
    for frag in fragments:
        add_textbox(slide, 1.02, y, 1.95, 0.18, frag, font_name="Inter", font_size=8.7, color=TEXT_SOFT)
        y += 0.28

    transform = add_box(slide, 3.76, 2.14, 2.32, 1.14, fill=ACCENT_LIGHT, line=ACCENT, line_width=1.0)
    add_textbox(slide, 4.02, 2.45, 1.78, 0.28, "Translate -> structure -> verify", font_name="Montserrat", font_size=12.2, color=ACCENT_DARK, bold=True, align=PP_ALIGN.CENTER)
    add_box(slide, 4.06, 3.38, 1.74, 0.42, fill=SOFT_PANEL, line=ACCENT, line_width=0.9)
    add_textbox(slide, 4.18, 3.50, 1.52, 0.12, "Necessary work that does not have to consume visit time", font_name="Inter", font_size=7.9, color=ACCENT_DARK, align=PP_ALIGN.CENTER)

    output = add_box(slide, 6.62, 1.80, 2.42, 2.00, fill=SOFT_PANEL, line=LINE, line_width=1.0)
    add_panel_label(slide, 6.88, 2.03, "Decision-ready structure")
    rows = [
        ("Key positives / negatives", 2.36),
        ("Timeline", 2.78),
        ("Questions to verify", 3.20),
    ]
    for label, row_y in rows:
        add_box(slide, 6.88, row_y, 1.88, 0.28, fill=LIGHT_GRAY, line=LIGHT_GRAY, line_width=0.5, shape_type=MSO_AUTO_SHAPE_TYPE.RECTANGLE)
        add_textbox(slide, 7.02, row_y + 0.06, 1.55, 0.12, label, font_name="Inter", font_size=8.3, color=TEXT_SOFT)

    add_line(slide, 3.33, 2.72, 3.76, 2.72, color=LINE, width=1.0)
    add_line(slide, 6.08, 2.72, 6.62, 2.72, color=ACCENT, width=1.0)

    add_footer(slide, footer_bird)


def build_slide_5(slide, footer_bird: bytes) -> None:
    clear_slide(slide)
    add_header(
        slide,
        "History is the cleanest place to reclaim minutes.",
        "It is repeated, variable, and can begin before the scheduled encounter.",
    )

    band = add_box(slide, 2.18, 2.18, 5.10, 0.96, fill=ACCENT, line=ACCENT, line_width=1.0)
    add_textbox(slide, 3.06, 2.47, 3.40, 0.22, "History reconstruction", font_name="Montserrat", font_size=15.0, color=RGBColor(0xFF, 0xFF, 0xFF), bold=True, align=PP_ALIGN.CENTER)

    callouts = [
        (0.82, 1.48, 2.10, 0.82, "Retold at every visit", "The same story gets repeated across encounters"),
        (0.82, 3.08, 2.10, 0.86, "Time-variable", "It takes different amounts of time depending on complexity and communication"),
        (7.34, 2.18, 1.82, 1.04, "Can begin before the encounter", "Much of it can happen before the clinician is on the clock"),
    ]
    for x, y, w, h, title, detail in callouts:
        add_box(slide, x, y, w, h, fill=SOFT_PANEL, line=LINE, line_width=0.95)
        add_textbox(slide, x + 0.16, y + 0.12, w - 0.32, 0.18, title, font_name="Montserrat", font_size=10.0, color=BLACK, bold=True)
        add_textbox(slide, x + 0.16, y + 0.34, w - 0.32, h - 0.34, detail, font_name="Inter", font_size=8.1, color=TEXT_SOFT)

    add_line(slide, 2.92, 2.18, 2.92, 1.96, color=LINE, width=0.9)
    add_line(slide, 2.92, 3.14, 2.92, 3.30, color=LINE, width=0.9)
    add_line(slide, 7.28, 2.66, 7.34, 2.66, color=ACCENT, width=1.0)

    add_footer(slide, footer_bird)


def build_slide_6(slide, footer_bird: bytes) -> None:
    clear_slide(slide)
    add_header(
        slide,
        "The hard part is producing a history clinicians can start from.",
        "Vox creates value by turning patient story into adaptive, structured, verifiable history clinicians can actually use.",
    )

    card = add_box(slide, 2.66, 1.72, 4.58, 1.28, fill=WARM_PANEL, line=ACCENT, line_width=1.1)
    add_textbox(slide, 3.18, 2.10, 3.55, 0.26, "History clinicians can start from", font_name="Montserrat", font_size=15.0, color=BLACK, bold=True, align=PP_ALIGN.CENTER)
    add_textbox(slide, 3.38, 2.42, 3.15, 0.18, "The standard Vox has to meet", font_name="Inter", font_size=8.9, color=TEXT_MUTED, align=PP_ALIGN.CENTER)

    support_titles = ["Adaptive", "Structured", "Prioritized", "Verifiable"]
    for idx, title in enumerate(support_titles):
        x = 1.38 + idx * 2.08
        add_box(slide, x, 3.30, 1.68, 0.62, fill=SOFT_PANEL, line=LINE, line_width=0.9)
        add_textbox(slide, x + 0.12, 3.50, 1.44, 0.14, title, font_name="Inter", font_size=9.2, color=TEXT_SOFT, bold=True, align=PP_ALIGN.CENTER)
    add_textbox(slide, 3.12, 4.12, 3.64, 0.18, "Changes the starting point of the visit", font_name="Inter", font_size=9.4, color=ACCENT_DARK, align=PP_ALIGN.CENTER, bold=True)

    add_footer(slide, footer_bird)


def build_slide_7(slide, footer_bird: bytes) -> None:
    clear_slide(slide)
    add_header(
        slide,
        "Vox prepares the visit before it starts.",
        "Vox interviews, structures, and hands off a brief the clinician can verify and use.",
    )

    steps = ["Invite", "Interview", "Structure", "Verify"]
    descriptions = [
        "Patient gets a pre-visit link",
        "Adaptive interview captures story",
        "History becomes a structured brief",
        "Clinician reviews, confirms, and starts",
    ]
    x_positions = [0.78, 2.95, 5.12, 7.29]
    for idx, (x, step, detail) in enumerate(zip(x_positions, steps, descriptions), start=1):
        add_box(slide, x, 2.10, 1.58, 1.24, fill=SOFT_PANEL, line=LINE, line_width=0.95)
        add_textbox(slide, x + 0.16, 2.34, 1.24, 0.20, f"{idx}. {step}", font_name="Montserrat", font_size=10.6, color=BLACK, bold=True, align=PP_ALIGN.CENTER)
        add_textbox(slide, x + 0.12, 2.66, 1.34, 0.42, detail, font_name="Inter", font_size=8.0, color=TEXT_SOFT, align=PP_ALIGN.CENTER)
        if idx < 4:
            add_line(slide, x + 1.58, 2.72, x + 2.17, 2.72, color=LINE, width=0.9)

    tags = [("provenance", 5.18), ("explicit unknowns", 6.08), ("verification queue", 7.29)]
    for label, x in tags:
        add_chip(slide, x, 3.56, 0.82 if label == "provenance" else 1.10 if label == "explicit unknowns" else 1.18, 0.26, label)

    add_footer(slide, footer_bird)


def build_slide_8(slide, footer_bird: bytes) -> None:
    clear_slide(slide)
    add_header(slide, "The clinician starts from a structured brief.", "Scan, verify, decide.")

    artifact = add_box(slide, 3.00, 1.58, 3.30, 2.84, fill=SOFT_PANEL, line=LINE, line_width=1.0, shape_type=MSO_AUTO_SHAPE_TYPE.RECTANGLE)
    add_panel_label(slide, 3.24, 1.86, "Signal Report artifact")
    sections = [
        ("Chief concern + context", 2.18, 0.58),
        ("Key positives / negatives", 2.84, 0.52),
        ("Timeline anchors", 3.44, 0.44),
        ("Questions to verify", 3.96, 0.52),
    ]
    for label, y, h in sections:
        add_box(slide, 3.24, y, 2.82, h, fill=LIGHT_GRAY, line=LIGHT_GRAY, line_width=0.5, shape_type=MSO_AUTO_SHAPE_TYPE.RECTANGLE)
        add_textbox(slide, 3.38, y + 0.08, 2.4, 0.12, label, font_name="Inter", font_size=8.1, color=TEXT_SOFT)

    callouts = [
        (0.76, 1.96, 1.88, 0.62, "What matters"),
        (6.82, 1.92, 2.02, 0.62, "What happened when"),
        (6.82, 3.18, 2.02, 0.62, "What still needs confirmation"),
    ]
    anchors = [(3.00, 2.28), (6.30, 2.28), (6.30, 3.72)]
    for (x, y, w, h, text), (ax, ay) in zip(callouts, anchors):
        add_box(slide, x, y, w, h, fill=WARM_PANEL, line=LINE, line_width=0.9)
        add_textbox(slide, x + 0.12, y + 0.21, w - 0.24, 0.18, text, font_name="Montserrat", font_size=9.2, color=BLACK, bold=True, align=PP_ALIGN.CENTER)
        add_line(slide, ax, ay, x + (w if x < 3 else 0), y + 0.31, color=LINE, width=0.8)

    add_footer(slide, footer_bird)


def build_slide_9(slide, footer_bird: bytes) -> None:
    clear_slide(slide)
    add_header(
        slide,
        "The visit still starts from zero.",
        "Clinicians have shown they will adopt assistive tools. The open problem is starting the visit with better structure.",
    )

    zones = [
        (0.86, 2.08, 2.42, 1.48, "Before visit", WARM_PANEL, LINE, ["invite", "patient context", "prep signals"]),
        (3.46, 2.08, 2.42, 1.48, "Start of visit", RGBColor(0xFF, 0xFF, 0xFF), ACCENT, ["structured starting point", "pre-decision signal", "recovered decision time"]),
        (6.06, 2.08, 2.42, 1.48, "During / after", LIGHT_GRAY, LINE, ["note generation", "clerical relief", "documentation support"]),
    ]
    for x, y, w, h, title, fill, border, lines in zones:
        add_box(slide, x, y, w, h, fill=fill, line=border, line_width=1.0)
        add_textbox(slide, x + 0.16, y + 0.18, w - 0.32, 0.18, title, font_name="Montserrat", font_size=10.0, color=BLACK, bold=True)
        row_y = y + 0.54
        for line_text in lines:
            add_textbox(slide, x + 0.18, row_y, w - 0.36, 0.14, line_text, font_name="Inter", font_size=8.0, color=TEXT_SOFT)
            row_y += 0.24

    add_footer(slide, footer_bird)


def build_slide_10(slide, footer_bird: bytes) -> None:
    clear_slide(slide)
    add_header(
        slide,
        "Recovered visit time creates real capacity.",
        "When recovered minutes become completed visits, the impact is meaningful.",
    )

    hero = add_box(slide, 0.78, 1.72, 3.44, 2.42, fill=WARM_PANEL, line=ACCENT, line_width=1.1)
    add_textbox(slide, 1.04, 2.00, 2.9, 0.20, "+1 visit / provider / day", font_name="Montserrat", font_size=11.6, color=BLACK, bold=True)
    add_textbox(slide, 1.04, 2.34, 2.9, 0.52, "220 clinic days\n$125 / visit\n40% contribution margin", font_name="Inter", font_size=8.1, color=TEXT_SOFT)
    add_textbox(slide, 1.04, 3.02, 2.9, 0.38, "$11,000 / year", font_name="Inter", font_size=24.0, color=BLACK, bold=True)
    add_textbox(slide, 1.04, 3.52, 2.9, 0.18, "~$917 / month break-even spend / provider", font_name="Inter", font_size=8.4, color=TEXT_MUTED)

    scenarios = [
        ("Conservative", "$7,920 / year", "+1 / day at $90 / visit"),
        ("Base", "$11,000 / year", "+1 / day at $125 / visit"),
        ("Upside", "$22,000 / year", "+2 / day at $125 / visit"),
    ]
    for idx, (label, amount, detail) in enumerate(scenarios):
        y = 1.74 + idx * 0.82
        add_box(slide, 4.82, y, 4.05, 0.66, fill=SOFT_PANEL, line=LINE, line_width=0.9)
        add_textbox(slide, 5.06, y + 0.10, 1.1, 0.16, label, font_name="Montserrat", font_size=8.8, color=BLACK, bold=True)
        add_textbox(slide, 6.00, y + 0.08, 1.6, 0.16, amount, font_name="Inter", font_size=10.8, color=BLACK, bold=True)
        add_textbox(slide, 7.18, y + 0.10, 1.40, 0.16, detail, font_name="Inter", font_size=8.0, color=TEXT_SOFT)

    add_textbox(slide, 4.90, 4.30, 3.70, 0.18, "Recovered time matters when it becomes completed visits", font_name="Inter", font_size=9.0, color=ACCENT_DARK, bold=True, align=PP_ALIGN.CENTER)

    add_footer(slide, footer_bird)


def build_slide_11(slide, footer_bird: bytes) -> None:
    clear_slide(slide)
    add_header(
        slide,
        "Start where throughput pressure is high and friction is low.",
        "High-volume settings, few decision makers, low integration burden.",
    )

    settings = [
        (0.82, 1.72, "Urgent care"),
        (0.82, 2.28, "Independent primary care groups"),
        (0.82, 2.84, "Multi-specialty groups"),
        (0.82, 3.40, "Health systems"),
        (0.82, 3.96, "FQHCs"),
    ]
    for x, y, label in settings:
        add_box(slide, x, y, 1.96, 0.34, fill=LIGHT_GRAY, line=LIGHT_GRAY, line_width=0.5)
        add_textbox(slide, x + 0.12, y + 0.09, 1.68, 0.12, label, font_name="Inter", font_size=8.2, color=TEXT_SOFT)

    filters = ["throughput pressure", "few decision makers", "low integration friction"]
    for idx, label in enumerate(filters):
        x = 3.50
        y = 1.94 + idx * 0.72
        add_box(slide, x, y, 1.70, 0.42, fill=WARM_PANEL, line=LINE, line_width=0.9)
        add_textbox(slide, x + 0.14, y + 0.12, 1.42, 0.14, label, font_name="Inter", font_size=8.3, color=BLACK, align=PP_ALIGN.CENTER)
        add_line(slide, 2.84, y + 0.21, x, y + 0.21, color=LINE, width=0.8)

    beachhead = add_box(slide, 6.00, 2.22, 2.30, 0.86, fill=ACCENT_LIGHT, line=ACCENT, line_width=1.1)
    add_textbox(slide, 6.20, 2.50, 1.90, 0.18, "Independent primary care groups", font_name="Montserrat", font_size=10.1, color=ACCENT_DARK, bold=True, align=PP_ALIGN.CENTER)
    add_line(slide, 5.20, 2.15, 6.00, 2.65, color=ACCENT, width=0.9)
    add_line(slide, 5.20, 2.87, 6.00, 2.65, color=ACCENT, width=0.9)

    proof = add_box(slide, 6.02, 3.42, 2.96, 1.08, fill=SOFT_PANEL, line=LINE, line_width=0.9)
    add_textbox(slide, 6.22, 3.60, 2.40, 0.18, "What the pilot must prove", font_name="Montserrat", font_size=9.3, color=BLACK, bold=True)
    add_textbox(
        slide,
        6.22,
        3.88,
        2.42,
        0.44,
        "completion rate\nclinician verification time\nre-asking reduction\ncaptured slots / schedule effect",
        font_name="Inter",
        font_size=8.0,
        color=TEXT_SOFT,
    )

    add_footer(slide, footer_bird)


def build_slide_12_team(slide, footer_bird: bytes, headshots: list[bytes]) -> None:
    clear_slide(slide)
    add_centered_header(
        slide,
        "Built by people who understand the visit.",
        "Clinical insight, product judgment, and deployment realism.",
    )

    founders = [
        ("Yashar Niknafs, MD, PhD", "Founder - CEO", "AI / product execution"),
        ("Balaji Pandian, MD, MBA", "Founder", "operations / commercialization"),
        ("Rana Kabeer, MD, MPH", "Founder", "clinical credibility"),
        ("Kristian Black, MD, MS", "Founder", "workflow / systems implementation"),
    ]
    x_positions = [0.44, 2.88, 5.20, 7.34]

    for idx, ((name, role, line), x) in enumerate(zip(founders, x_positions)):
        add_box(slide, x, 1.72, 1.92, 1.64, fill=SOFT_PANEL, line=LINE, line_width=0.8)
        slide.shapes.add_picture(BytesIO(headshots[idx]), emu(x + 0.30), emu(1.72), width=emu(1.33), height=emu(1.33))
        add_textbox(slide, x, 3.34, 1.92, 0.16, role, font_name="Montserrat", font_size=8.2, color=BLACK, align=PP_ALIGN.CENTER)
        add_textbox(slide, x, 3.56, 1.92, 0.18, name, font_name="Montserrat", font_size=8.2, color=BLACK, align=PP_ALIGN.CENTER)
        add_textbox(slide, x + 0.08, 3.84, 1.76, 0.26, line, font_name="Inter", font_size=7.8, color=TEXT_MUTED, align=PP_ALIGN.CENTER)

    add_footer(slide, footer_bird)


def build_slide_13_milestones(slide, footer_bird: bytes) -> None:
    clear_slide(slide)
    add_header(
        slide,
        "This round is about proving trust, capture, and repeatable deployment.",
        "Fund the milestones that turn product promise into deployment proof.",
    )

    milestones = [
        ("1", "Trust primitives fully working", "clinician-grade output\nauditability\nverification loop"),
        ("2", "Pilot sites live with measurable capture", "completion rate\nverification time\nre-asking reduction"),
        ("3", "Paid conversions + repeatable playbook", "initial paid sites\ndeployment motion\nrepeatable rollout"),
    ]
    for idx, (num, title, detail) in enumerate(milestones):
        x = 0.74 + idx * 3.04
        add_box(slide, x, 1.78, 2.52, 1.96, fill=WARM_PANEL if idx == 1 else SOFT_PANEL, line=ACCENT if idx == 1 else LINE, line_width=1.0)
        add_textbox(slide, x + 0.16, 1.98, 0.30, 0.20, num, font_name="Inter", font_size=12.0, color=ACCENT_DARK if idx == 1 else TEXT_SOFT, bold=True)
        add_textbox(slide, x + 0.16, 2.24, 2.10, 0.40, title, font_name="Montserrat", font_size=10.4, color=BLACK, bold=True)
        add_textbox(slide, x + 0.16, 2.82, 2.05, 0.54, detail, font_name="Inter", font_size=8.0, color=TEXT_SOFT)

    add_box(slide, 0.74, 4.12, 8.12, 0.42, fill=SOFT_PANEL, line=LINE, line_width=0.8, shape_type=MSO_AUTO_SHAPE_TYPE.RECTANGLE)
    add_textbox(slide, 0.96, 4.24, 1.8, 0.12, "Amount: $1M", font_name="Inter", font_size=8.2, color=TEXT_SOFT, bold=True)
    add_textbox(slide, 3.56, 4.24, 2.0, 0.12, "Timing: current round", font_name="Inter", font_size=8.2, color=TEXT_SOFT, bold=True)
    add_textbox(slide, 6.02, 4.24, 2.4, 0.12, "Use: trust primitives, pilots, deployment proof", font_name="Inter", font_size=8.2, color=TEXT_SOFT, bold=True)

    add_footer(slide, footer_bird)


def extract_assets(prs: Presentation) -> tuple[ImageAsset, bytes, list[bytes]]:
    slide1 = prs.slides[0]
    bird = slide1.shapes[0]
    title_bird = ImageAsset(
        blob=bird.image.blob,
        left=bird.left,
        top=bird.top,
        width=bird.width,
        height=bird.height,
    )

    footer_slide = prs.slides[10]
    footer_bird = footer_slide.shapes[13].image.blob

    team_slide = prs.slides[12]
    photo_shapes = []
    for shape in team_slide.shapes:
        if getattr(shape, "image", None) is not None and shape.width > emu(1.2) and shape.top < emu(3.0):
            photo_shapes.append(shape)
    photo_shapes.sort(key=lambda s: s.left)
    headshots = [shape.image.blob for shape in photo_shapes]

    return title_bird, footer_bird, headshots


def build_deck(source_deck: Path, output_deck: Path) -> Path:
    output_deck.parent.mkdir(parents=True, exist_ok=True)
    prs = Presentation(str(source_deck))
    delete_slides_after(prs, 13)
    title_bird, footer_bird, headshots = extract_assets(prs)

    build_slide_1(prs.slides[0], footer_bird)
    build_slide_2(prs.slides[1], footer_bird)
    build_slide_3(prs.slides[2], footer_bird)
    build_slide_4(prs.slides[3], footer_bird)
    build_slide_5(prs.slides[4], footer_bird)
    build_slide_6(prs.slides[5], footer_bird)
    build_slide_7(prs.slides[6], footer_bird)
    build_slide_8(prs.slides[7], footer_bird)
    build_slide_9(prs.slides[8], footer_bird)
    build_slide_10(prs.slides[9], footer_bird)
    build_slide_11(prs.slides[10], footer_bird)
    build_slide_12_team(prs.slides[11], footer_bird, headshots)
    build_slide_13_milestones(prs.slides[12], footer_bird)

    prs.save(str(output_deck))
    return output_deck


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Rebuild the Vox pitch-deck outline stub from a source PPTX."
    )
    parser.add_argument(
        "--source",
        type=Path,
        default=DEFAULT_SOURCE_DECK,
        help=f"Source PPTX path (default: {DEFAULT_SOURCE_DECK})",
    )
    parser.add_argument(
        "--output",
        type=Path,
        default=DEFAULT_OUTPUT_DECK,
        help=f"Output PPTX path (default: {DEFAULT_OUTPUT_DECK})",
    )
    return parser.parse_args()


if __name__ == "__main__":
    args = parse_args()
    path = build_deck(args.source, args.output)
    print(path)
